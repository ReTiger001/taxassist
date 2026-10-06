"""附件解析测试。

关键约定：**扫描件（no_text_layer）不是失败**。
它和"解析器坏了（failed）"必须区分——前者人打开看一眼就行，后者要改代码。
"""
from __future__ import annotations

import pytest

from taxassist.collect.attachments import (
    parse_attachment,
    safe_filename,
)


@pytest.fixture()
def pdf_with_text(tmp_path):
    pymupdf = pytest.importorskip("pymupdf")
    path = tmp_path / "policy.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Policy text layer sample")
    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture()
def pdf_without_text(tmp_path):
    """模拟扫描件：有页面但无文本层。"""
    pymupdf = pytest.importorskip("pymupdf")
    path = tmp_path / "scan.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture()
def xlsx_file(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    path = tmp_path / "form.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "申报表"
    ws.append(["项目", "金额"])
    ws.append(["营业收入", 1000])
    wb.save(str(path))
    return path


def test_pdf_text_extracted(pdf_with_text):
    text, status = parse_attachment(pdf_with_text)
    assert status == "ok"
    assert "Policy text layer sample" in text


def test_sniff_kind_reads_the_real_format(tmp_path):
    from taxassist.collect.attachments import sniff_kind

    p = tmp_path / "a.bin"
    p.write_bytes(b"%PDF-1.7\nrest")
    assert sniff_kind(p) == "pdf"
    p.write_bytes(b"PK\x03\x04rest")
    assert sniff_kind(p) == "zip"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest")
    assert sniff_kind(p) == "ole2"
    p.write_bytes(b"plain text")
    assert sniff_kind(p) is None


def test_ole2_file_named_docx_is_unsupported_not_failed(tmp_path):
    """实测踩过：税务局把老式 .doc 挂成 .docx，文件头是 d0cf11e0。
    按扩展名交给 python-docx 会抛 PackageNotFoundError，于是被记成 failed ——
    可文件是好的，是我们的工具选错了。形式上是"我们读不了"，
    就该如实标 unsupported，别让人以为数据坏了。"""
    from taxassist.collect.attachments import parse_attachment

    p = tmp_path / "伪装.docx"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    text, status = parse_attachment(p)
    assert text is None
    assert status == "unsupported"


def test_zip_file_named_doc_is_parsed_as_docx(tmp_path):
    """反向不符：扩展名说老式 .doc，内容其实是 ZIP —— 按内容走。"""
    import docx

    from taxassist.collect.attachments import parse_attachment

    p = tmp_path / "其实是新式.doc"
    d = docx.Document()
    d.add_paragraph("增值税加计抵减政策")
    d.save(str(p))
    text, status = parse_attachment(p)
    assert status == "ok"
    assert "增值税加计抵减政策" in text


def test_missing_file_reports_failed_not_crash(tmp_path):
    from taxassist.collect.attachments import parse_attachment

    text, status = parse_attachment(tmp_path / "没有这个文件.docx")
    assert text is None
    assert status.startswith("failed:")


def test_scanned_pdf_is_not_a_failure(pdf_without_text):
    text, status = parse_attachment(pdf_without_text)
    assert status == "no_text_layer"      # 不是 "failed"
    assert text is None


def test_xlsx_extracted_with_sheet_name(xlsx_file):
    text, status = parse_attachment(xlsx_file)
    assert status == "ok"
    assert "[工作表] 申报表" in text
    assert "营业收入" in text


def test_unsupported_format_is_labelled(tmp_path):
    p = tmp_path / "old.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0fake")
    _, status = parse_attachment(p)
    assert status == "unsupported"


def test_corrupt_file_reports_failure_not_silence(tmp_path):
    p = tmp_path / "broken.pdf"
    p.write_bytes(b"this is not a pdf at all")
    _, status = parse_attachment(p)
    assert status.startswith("failed:")


def test_parse_docx_falls_back_to_xml(tmp_path):
    """python-docx 打不开的 docx，也要能读出文字。

    实测（2026-10-06）：6 条税务局的 .docx 让 python-docx 抛 KeyError
    （文件里缺它期望的某个 XML 部件，但 Word/WPS 能正常打开）。我们只要
    文本，直读 word/document.xml 就够 —— 这类文件不该判成 failed。

    这里故意只写 word/document.xml、不写 python-docx 需要的那套部件，
    模拟那条真实路径。
    """
    import zipfile

    from taxassist.collect.attachments import parse_docx

    p = tmp_path / "no_parts.docx"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr(
            "word/document.xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org'
            '/wordprocessingml/2006/main"><w:body>'
            "<w:p><w:r><w:t>第一段文字</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>第二段文字</w:t></w:r></w:p>"
            "</w:body></w:document>")

    text = parse_docx(p)
    assert "第一段文字" in text
    assert "第二段文字" in text


def test_parse_docx_missing_part_returns_empty(tmp_path):
    """连 word/document.xml 都没有的 zip 不该抛异常，返回空串即可。

    空串会让 _run_parser 记成 no_text_layer（"人打开看一眼就行"），
    而不是 failed（"要改代码"）—— 这个区分是本模块的核心约定。
    """
    import zipfile

    from taxassist.collect.attachments import parse_docx

    p = tmp_path / "empty.docx"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr("dummy.txt", "not a docx")
    assert parse_docx(p) == ""


def test_safe_filename_strips_dangerous_chars():
    assert safe_filename("《申报表》及其附列资料.xls") == "《申报表》及其附列资料.xls"
    assert "/" not in safe_filename("a/b\\c:d.xls")
    assert safe_filename("") == "attachment"
    assert len(safe_filename("长" * 300)) <= 120
