"""汇算清缴计算与底稿测试。

重点验证三件事：
1. 限额计算的**算式**对不对（业务招待费的双限额最容易错）
2. 依据核查能否识别"库中已有但已废止"——这是最该报警的状态
3. 底稿里**必须**出现"未考虑事项"，否则会被误当成完整申报表
"""
from __future__ import annotations

import pytest

from taxassist import db as dbmod
from taxassist.cit import engine, workpaper
from taxassist.cit.engine import BASIS_OK, BASIS_STALE, CitInputs


@pytest.fixture()
def conn(tmp_path):
    c = dbmod.connect(tmp_path / "t.db")
    dbmod.init_db(c)
    yield c
    c.close()


# ---------------------------------------------------------------- 计算

def test_entertainment_uses_lower_of_two_limits():
    """业务招待费：发生额 60% 与营业收入 5‰ 孰低。

    收入 10,000,000 → 5‰ = 50,000；发生额 100,000 → 60% = 60,000。
    孰低 50,000 → 调增 50,000。
    """
    r = engine.compute(CitInputs(operating_revenue=10_000_000, entertainment=100_000,
                                 total_profit=1_000_000))
    item = next(i for i in r.items if i.code == "entertainment")
    assert item.tax_amount == pytest.approx(50_000)
    assert item.increase == pytest.approx(50_000)
    assert item.carryforward == 0        # 业务招待费超限不得结转
    assert "5‰" in item.detail or "0.50%" in item.detail or "50,000" in item.detail


def test_entertainment_limited_by_occurrence_ratio_when_revenue_is_large():
    """收入很大时，60% 那一档成为限制。"""
    r = engine.compute(CitInputs(operating_revenue=100_000_000, entertainment=100_000))
    item = next(i for i in r.items if i.code == "entertainment")
    assert item.tax_amount == pytest.approx(60_000)
    assert item.increase == pytest.approx(40_000)


def test_welfare_education_union_caps():
    r = engine.compute(CitInputs(wages=1_000_000, welfare=200_000,
                                 education=100_000, union=30_000))
    by = {i.code: i for i in r.items}
    assert by["welfare"].increase == pytest.approx(200_000 - 140_000)
    assert by["education"].increase == pytest.approx(100_000 - 80_000)
    assert by["education"].carryforward == pytest.approx(20_000)   # 教育经费可结转
    assert by["union"].increase == pytest.approx(30_000 - 20_000)


def test_donation_cap_and_carryforward():
    r = engine.compute(CitInputs(total_profit=1_000_000, donation=150_000))
    item = next(i for i in r.items if i.code == "donation")
    assert item.tax_amount == pytest.approx(120_000)
    assert item.increase == pytest.approx(30_000)
    assert item.carryforward == pytest.approx(30_000)


def test_disallow_items_increase_full_amount():
    r = engine.compute(CitInputs(fines=50_000, provisions=80_000))
    by = {i.code: i for i in r.items}
    assert by["fines"].increase == pytest.approx(50_000)
    assert by["provisions"].increase == pytest.approx(80_000)
    assert by["fines"].decrease == 0


def test_rd_super_deduction_reduces_taxable_income():
    r = engine.compute(CitInputs(rd_expense=1_000_000))
    item = next(i for i in r.items if i.code == "rd_expense")
    assert item.decrease == pytest.approx(1_000_000)   # 100% 加计
    assert item.increase == 0


def test_totals_and_tax_payable():
    r = engine.compute(CitInputs(
        total_profit=1_000_000, operating_revenue=10_000_000,
        entertainment=100_000, fines=50_000, rd_expense=200_000))
    expected_inc = r.total_increase
    expected_dec = r.total_decrease
    assert r.taxable_income == pytest.approx(1_000_000 + expected_inc - expected_dec)
    assert r.tax_payable == pytest.approx(r.taxable_income * 0.25)


def test_zero_inputs_produce_no_line_items():
    r = engine.compute(CitInputs())
    assert r.items == []
    assert r.taxable_income == 0


def test_extra_manual_adjustment_is_recorded():
    r = engine.compute(CitInputs(total_profit=100_000, extra_increase=5_000,
                                 extra_note="补提折旧差异"))
    extra = next(i for i in r.items if i.code == "extra")
    assert extra.increase == pytest.approx(5_000)
    assert "补提折旧" in extra.detail


# ---------------------------------------------------------------- 依据核查

def test_basis_marks_ok_when_library_has_valid_doc(conn):
    from taxassist import store
    store.upsert_policy(conn, {
        "doc_uid": "d1", "title": "国家税务总局关于研发费用加计扣除政策的公告",
        "cwrq": "2026-01-01", "p_effect_status": "现行有效", "p_effect_source": "official",
    })
    rule = next(r for r in engine.RULES if r.code == "rd_expense")
    status, evidence = engine.check_basis(conn, rule)
    assert status == BASIS_OK
    assert "研发费用" in evidence


def test_basis_marks_stale_when_library_doc_repealed(conn):
    """依据已废止是最该报警的状态 —— 拿着废止文件算出来的底稿是有害的。"""
    from taxassist import store
    store.upsert_policy(conn, {
        "doc_uid": "d2", "title": "关于业务招待费扣除标准的旧规定",
        "cwrq": "2010-01-01", "p_effect_status": "已废止", "p_effect_source": "inferred",
    })
    rule = next(r for r in engine.RULES if r.code == "entertainment")
    status, _ = engine.check_basis(conn, rule)
    assert status == BASIS_STALE


def test_basis_unknown_when_not_in_library(conn):
    rule = next(r for r in engine.RULES if r.code == "union")
    status, _ = engine.check_basis(conn, rule)
    assert status in (engine.BASIS_UNKNOWN, engine.BASIS_NO_DB)


def test_basis_without_conn_is_not_checked():
    rule = next(r for r in engine.RULES if r.code == "welfare")
    status, _ = engine.check_basis(None, rule)
    assert status == engine.BASIS_NO_DB


# ---------------------------------------------------------------- 底稿

def test_warnings_always_present():
    """底稿必须自带"未考虑事项"，否则会被当成完整申报表。"""
    r = engine.compute(CitInputs(total_profit=100))
    assert r.warnings
    joined = "".join(r.warnings)
    for must in ("弥补", "优惠", "境外"):
        assert must in joined


def test_workpaper_roundtrip(tmp_path):
    tpl = workpaper.write_input_template(tmp_path / "tpl.xlsx")
    assert tpl.exists()

    from openpyxl import load_workbook
    wb = load_workbook(str(tpl))
    ws = wb["输入数据"]
    # 按标签填两个值
    for row in ws.iter_rows(min_row=5):
        if row[0].value == "营业收入":
            row[1].value = 10_000_000
        if row[0].value == "业务招待费发生额":
            row[1].value = 100_000
    wb.save(str(tpl))
    wb.close()

    inputs = workpaper.read_inputs(tpl)
    assert inputs.operating_revenue == pytest.approx(10_000_000)
    assert inputs.entertainment == pytest.approx(100_000)

    out = workpaper.export_workpaper(engine.compute(inputs), tmp_path / "wp.xlsx")
    assert out.exists()
    wb2 = load_workbook(str(out))
    assert set(wb2.sheetnames) >= {"调整明细", "汇总", "输入数据", "未考虑事项"}
    # 未考虑事项页非空
    assert wb2["未考虑事项"]["A1"].value
    wb2.close()


def test_workpaper_carries_basis_status(conn, tmp_path):
    from taxassist import store
    store.upsert_policy(conn, {
        "doc_uid": "d3", "title": "关于职工福利费的旧规定",
        "cwrq": "2010-01-01", "p_effect_status": "已废止", "p_effect_source": "inferred",
    })
    result = engine.compute(CitInputs(wages=100_000, welfare=50_000), conn=conn)
    item = next(i for i in result.items if i.code == "welfare")
    assert item.basis_status == BASIS_STALE

    out = workpaper.export_workpaper(result, tmp_path / "wp2.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(str(out))
    ws = wb["调整明细"]
    statuses = [ws.cell(row=r, column=9).value for r in range(5, ws.max_row + 1)]
    assert any(s and s.startswith("库中已收录·已废止") for s in statuses)
    wb.close()


# ---------------------------------------------------------------- 结构性约束

def test_every_rule_code_matches_an_input_field():
    """规则的 code 必须与 CitInputs 的字段名一致。

    引擎按 ``code`` 去输入对象里取值；不一致时该项**永远取到 0**，
    静默算出一份少了调整项的底稿，而底稿看起来完全正常。
    实测踩过：研发费用加计扣除的 code 曾写成 rd_super_deduction，导致该项从未生效。
    """
    fields = set(CitInputs.__dataclass_fields__)
    for rule in engine.RULES:
        assert rule.code in fields, (
            f"规则 {rule.name} 的 code={rule.code!r} 不是 CitInputs 的字段，"
            f"该项将永远取到 0。可用字段：{sorted(fields)}")


def test_every_rule_declares_basis_and_formula():
    """每条规则都必须写明依据与算式 —— 底稿上要给出可核对的过程。"""
    for rule in engine.RULES:
        assert rule.basis, rule.code
        assert rule.formula, rule.code
        if rule.kind == "cap":
            assert rule.limit_ratio is not None and rule.limit_base, rule.code
            assert rule.limit_base in CitInputs.__dataclass_fields__, rule.code
